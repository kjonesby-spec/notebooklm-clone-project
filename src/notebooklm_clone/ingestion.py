from pathlib import Path
from urllib.parse import urlparse

import requests
from bs4 import BeautifulSoup
from pptx import Presentation
from pypdf import PdfReader

from .text_utils import clean_text


SUPPORTED_EXTENSIONS = {".pdf", ".pptx", ".txt", ".md"}


def extract_file(path: str) -> tuple[str, str]:
    file_path = Path(path)
    suffix = file_path.suffix.lower()
    if suffix not in SUPPORTED_EXTENSIONS:
        raise ValueError(f"Unsupported file type: {suffix}. Use PDF, PPTX, TXT, or MD.")
    if suffix == ".pdf":
        text = "\n\n".join(page.extract_text() or "" for page in PdfReader(file_path).pages)
    elif suffix == ".pptx":
        presentation = Presentation(file_path)
        slides = []
        for index, slide in enumerate(presentation.slides, start=1):
            content = [shape.text for shape in slide.shapes if hasattr(shape, "text")]
            slides.append(f"Slide {index}\n" + "\n".join(content))
        text = "\n\n".join(slides)
    else:
        text = file_path.read_text(encoding="utf-8", errors="replace")
    text = clean_text(text)
    if not text:
        raise ValueError(f"No readable text was found in {file_path.name}")
    return file_path.name, text


def extract_url(url: str) -> tuple[str, str]:
    parsed = urlparse(url.strip())
    if parsed.scheme not in {"http", "https"} or not parsed.netloc:
        raise ValueError("Enter a complete http:// or https:// URL")
    response = requests.get(url, timeout=20, headers={"User-Agent": "NotebookLMClone/1.0"})
    response.raise_for_status()
    soup = BeautifulSoup(response.text, "html.parser")
    for element in soup(["script", "style", "nav", "footer", "form"]):
        element.decompose()
    title = clean_text(soup.title.get_text(" ") if soup.title else parsed.netloc)
    text = clean_text(soup.get_text("\n"))
    if len(text) < 50:
        raise ValueError("The webpage did not provide enough readable text")
    return title, text

